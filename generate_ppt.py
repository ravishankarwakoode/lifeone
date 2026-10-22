from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create PowerPoint presentation
prs = Presentation()
prs.core_properties.title = "LifeOne - Daily Life Routing App"
prs.core_properties.subject = "Project presentation"
prs.core_properties.author = "LifeOne Team"

# Slide size: widescreen
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
NAVY = RGBColor(17, 24, 39)
BLUE = RGBColor(79, 70, 229)
GREEN = RGBColor(16, 185, 129)
GREY = RGBColor(107, 114, 128)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(243, 244, 246)


def add_title_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT

    title_box = slide.shapes.title
    title_box.text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(28)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = NAVY

    if subtitle:
        subtitle_box = slide.placeholders[1]
        subtitle_box.text = subtitle
        subtitle_box.text_frame.paragraphs[0].font.size = Pt(16)
        subtitle_box.text_frame.paragraphs[0].font.color.rgb = GREY


def add_content_slide(title, bullets, theme_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    top_bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.5))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = theme_color
    top_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.75), Inches(10), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = NAVY

    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.6), Inches(5.2))
    tf = body_box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.font.size = Pt(18)
        p.font.color.rgb = NAVY
        p.space_after = Pt(10)


def add_two_column_slide(title, left_items, right_items, left_color=BLUE, right_color=GREEN):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    top_bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.5))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = BLUE
    top_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(12), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = NAVY

    left_box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(5.1), Inches(4.8))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = left_color
    left_box.line.fill.background()

    right_box = slide.shapes.add_shape(1, Inches(6.2), Inches(1.5), Inches(5.9), Inches(4.8))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = right_color
    right_box.line.fill.background()

    left_tf = left_box.text_frame
    left_tf.word_wrap = True
    left_tf.margin_left = Inches(0.2)
    left_tf.margin_top = Inches(0.2)
    left_tf.margin_right = Inches(0.2)
    left_tf.margin_bottom = Inches(0.2)
    for item in left_items:
        p = left_tf.paragraphs[0] if left_tf.paragraphs[0].text == "" else left_tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(17)
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)

    right_tf = right_box.text_frame
    right_tf.word_wrap = True
    right_tf.margin_left = Inches(0.2)
    right_tf.margin_top = Inches(0.2)
    right_tf.margin_right = Inches(0.2)
    right_tf.margin_bottom = Inches(0.2)
    for item in right_items:
        p = right_tf.paragraphs[0] if right_tf.paragraphs[0].text == "" else right_tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(17)
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)


# Slide 1
add_title_slide("LifeOne", "Daily Life Routing App\nSmart daily planning, route management, and routine tracking")

# Slide 2: Overview
add_content_slide(
    "Project Overview",
    [
        "LifeOne is a mobile application for daily life routing and task management.",
        "It helps users organize routine activities, route schedules, and reminders.",
        "The system focuses on simplicity, time-saving, and daily productivity.",
        "The app is built with React Native and Expo for cross-platform support."
    ]
)

# Slide 3: Problem Statement
add_content_slide(
    "Problem Statement",
    [
        "People often manage tasks, travel, school runs, and work schedules separately.",
        "Daily routines become hard to track without a unified plan.",
        "Delay in planning can cause missed appointments and confusion.",
        "A simple daily route planner can reduce stress and improve productivity."
    ],
    theme_color=GREEN
)

# Slide 4: Objectives
add_content_slide(
    "Objectives",
    [
        "Provide login and registration for users.",
        "Display daily route and task timeline.",
        "Offer route planner for commute and travel suggestions.",
        "Enable reminders, checklist tracking, and profile customization.",
        "Make daily planning smart, easy, and user-friendly."
    ],
    theme_color=BLUE
)

# Slide 5: User Flow
add_content_slide(
    "User Flow",
    [
        "User opens the app.",
        "User logs in or creates an account.",
        "User sees Home and daily tasks.",
        "User checks Dashboard, Planner, Checklist, and Calendar.",
        "User reviews notifications, profile, and settings."
    ],
    theme_color=GREEN
)

# Slide 6: Key Screens
add_two_column_slide(
    "Authentication Screens",
    [
        "Login page with email and password fields.",
        "Register page with name, email, and password.",
        "Simple and clear user onboarding flow."
    ],
    [
        "Secure authentication experience.",
        "Users access app dashboard after login.",
        "Encourages easy daily usage."
    ],
    left_color=BLUE,
    right_color=GREEN
)

# Slide 7: Home Screen
add_content_slide(
    "Home Screen",
    [
        "Shows the daily route timeline.",
        "Displays wake-up routine, travel, work block, dinner, and errands.",
        "Each item includes time, location, and completion status.",
        "Helps users track the day at a glance."
    ],
    theme_color=BLUE
)

# Slide 8: Dashboard & Route Planner
add_content_slide(
    "Dashboard and Route Planner",
    [
        "Shows metrics such as trips today, minutes saved, and task completion.",
        "Lists upcoming trips for the day.",
        "Suggests route timing and travel conditions.",
        "Helps users plan efficient and stress-free movement."
    ],
    theme_color=GREEN
)

# Slide 9: Tasks, Calendar, and Route Details
add_content_slide(
    "Task and Schedule Management",
    [
        "Checklist page tracks daily tasks and completion status.",
        "Calendar page organizes planned activities and events.",
        "Add Task page enables new activity creation.",
        "Route Details screen displays start point, destination, time, and traffic status."
    ],
    theme_color=BLUE
)

# Slide 10: Notifications, Profile, Settings
add_content_slide(
    "Notifications, Profile, and Settings",
    [
        "Notifications remind users about route updates and reminders.",
        "Profile page stores user details and activity statistics.",
        "Settings page enables notifications, travel alerts, and dark mode preferences.",
        "This improves personalization and daily comfort."
    ],
    theme_color=GREEN
)

# Slide 11: Technology Stack
add_content_slide(
    "Technology Stack",
    [
        "React Native for mobile application development.",
        "Expo for rapid app setup and testing.",
        "React Navigation for screen flow and tab navigation.",
        "JavaScript for app logic and UI development."
    ],
    theme_color=BLUE
)

# Slide 12: Future Enhancements
add_content_slide(
    "Future Improvements",
    [
        "Real map integration with GPS and navigation.",
        "Firebase authentication and cloud storage.",
        "Push notifications for important reminders.",
        "AI-based route optimization and trip suggestions.",
        "Expanded analytics for user routines and habits."
    ],
    theme_color=GREEN
)

# Slide 13: Conclusion
add_content_slide(
    "Conclusion",
    [
        "LifeOne is designed for everyday organization and route planning.",
        "It combines task management, routing, reminders, and personal scheduling in one app.",
        "The application is practical, user-friendly, and ready for future growth."
    ],
    theme_color=BLUE
)

# Final slide
add_title_slide("Thank You", "LifeOne - Daily Life Routing App\nProject presentation complete")

# Save the presentation
output_file = "LifeOne_Project_Presentation.pptx"
prs.save(output_file)
print(f"PowerPoint generated successfully: {output_file}")
